from sqlalchemy import create_engine, MetaData, Table
from openpyxl import Workbook
from kolokwia import DB_FILE

EXP_FILE = 'DaneWrazliwe/kol.xlsx'

import pickle
from sqlalchemy import create_engine, MetaData
from openpyxl import Workbook

def export_database_to_excel(db_url, output_file):
    # Create SQLAlchemy engine and connect to the database
    engine = create_engine(db_url)
    connection = engine.connect()

    # Create metadata object to reflect database schema
    metadata = MetaData()
    metadata.reflect(bind=engine)

    # Create a workbook object
    wb = Workbook()

    # Loop through tables in the database
    for table_name, table in metadata.tables.items():
        # Create a new sheet for each table
        ws = wb.create_sheet(title=table_name)

        # Fetch all records from the table
        query = table.select()
        result = connection.execute(query)

        # Write column headers to the first row of the sheet
        headers = [column.name for column in table.columns]
        ws.append(headers)

        # Write data rows to the sheet
        for row in result:
            row_data = []
            for value in row:
                if isinstance(value, bytes):
                    # Convert bytes to string for PickleType/BLOB columns
                    try:
                        value = pickle.loads(value)
                        value = str(value)  # Convert to string
                    except pickle.UnpicklingError:
                        value = "<Unpicklable>"  # Handle unpicklable data
                row_data.append(value)
            ws.append(row_data)

        # Close the result set
        result.close()

    # Save the workbook to the output file
    wb.save(output_file)

    # Close database connection
    connection.close()

# Example usage:
export_database_to_excel('sqlite:///example.db', 'database_export.xlsx')


def main():
    export_database_to_excel(DB_FILE, EXP_FILE)

if __name__ == "__main__":
    main()
